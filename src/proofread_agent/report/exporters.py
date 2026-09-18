"""审校报告导出：Word(.docx) 与 PDF。

与 ``render_markdown`` 保持同构，便于用户在不同介质拿到一致内容。
PDF 使用 reportlab 内置的 STSong-Light CID 字体渲染中文，无需外部字体文件，
可直接在 Streamlit Cloud 等无中文字体环境中工作。
"""

from __future__ import annotations

import io
from typing import Any, Dict, List, Tuple

from .models import AuditReport

_GATE = {
    "BLOCK": ("#cf1322", "不予签发", "存在致命级问题，必须修改后重新审校"),
    "REVIEW": ("#d46b08", "退改", "存在需要修改的问题，建议修改后再审"),
    "PASS": ("#389e0d", "通过", "未发现阻断性问题"),
}


# ======================================================================
# 共享：从 report 抽取各区块数据
# ======================================================================
def _sections(report: AuditReport) -> Dict[str, Any]:
    d = report.to_dict()
    return {
        "gate": _GATE.get(report.release_gate, _GATE["PASS"]),
        "counts": report.counts,
        "dimensions": report.dimension_scores,
        "grouped": report.by_category(),
        "fact_check": d.get("fact_check") or [],
        "critic_notes": d.get("critic_notes") or [],
        "revision": (d.get("llm_notes") or {}).get("revision") or {},
        "trace": d.get("trace") or [],
    }


def _issue_lines(issue) -> List[str]:
    """单条问题的多行文本（与 Markdown 渲染一致）。"""
    lines = [f"[{issue.severity.label}] {issue.message}",
             f"位置：第 {issue.line} 行，第 {issue.paragraph + 1} 段"
             f"（字符 {issue.span.start}–{issue.span.end}）"]
    if issue.original:
        lines.append(f"原文：{issue.original}")
    if issue.suggestion:
        lines.append(f"建议：{issue.suggestion}")
    if issue.explain:
        lines.append(f"依据：{issue.explain.replace(chr(10), ' ')}")
    if issue.source:
        lines.append(f"出处：{issue.source}")
    lines.append(f"规则：{issue.rule_id}｜置信度 {issue.confidence}｜"
                 f"检出 Agent：{issue.agent}")
    return lines


# ======================================================================
# Word (.docx)
# ======================================================================
def export_docx(report: AuditReport, text: str = "") -> bytes:
    from docx import Document as Docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def _set_cjk(style, font_name: str = "宋体") -> None:
        style.font.name = font_name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), font_name)

    s = _sections(report)
    doc = Docx()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    _set_cjk(normal, "宋体")

    doc.add_heading(f"智能审校报告｜{report.doc_title}", level=0)
    gate_color, gate_label, gate_desc = s["gate"]
    p = doc.add_paragraph()
    r = p.add_run(f"签发门禁：{gate_label}（{report.release_gate}）")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(gate_color.lstrip("#"))
    p.add_run(f"　综合评分 {report.overall_score}/100　文体：{report.doc_type}"
              f"　风格预设：{report.style_preset or '未指定'}")
    doc.add_paragraph(
        f"字数 {report.char_count}｜段落 {report.paragraph_count}｜"
        f"术语库版本 {report.terminology_version}｜推理后端 {report.llm_backend}"
        f"{'（已降级为规则引擎）' if report.degraded else ''}")

    # 一、维度评分
    doc.add_heading("一、维度评分", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Light Grid Accent 1"
    for cell, t in zip(table.rows[0].cells,
                       ["维度", "得分", "致命", "严重", "一般", "提示", "合计"]):
        cell.text = t
    for ds in s["dimensions"]:
        cells = table.add_row().cells
        for cell, v in zip(cells, [ds.dimension.value, ds.score, ds.fatal,
                                    ds.major, ds.minor, ds.info, ds.total]):
            cell.text = str(v)
    cells = table.add_row().cells
    for cell, v in zip(cells, ["综合", report.overall_score, s["counts"]["FATAL"],
                                s["counts"]["MAJOR"], s["counts"]["MINOR"],
                                s["counts"]["INFO"], len(report.issues)]):
        cell.text = str(v)

    # 二、审校结论
    doc.add_heading("二、审校结论", level=1)
    doc.add_paragraph(report.summary or "（无）")

    # 三、问题清单
    doc.add_heading("三、问题清单", level=1)
    for cat, items in s["grouped"].items():
        doc.add_heading(f"{cat}（{len(items)} 条）", level=2)
        for idx, issue in enumerate(items, 1):
            ph = doc.add_paragraph()
            rr = ph.add_run(f"{idx}. {_issue_lines(issue)[0]}")
            rr.bold = True
            for line in _issue_lines(issue)[1:]:
                doc.add_paragraph(line, style="List Bullet")

    # 四、数据与事实复算
    if s["fact_check"]:
        doc.add_heading("四、数据与事实复算记录", level=1)
        ft = doc.add_table(rows=1, cols=5)
        ft.style = "Light Grid Accent 1"
        for cell, t in zip(ft.rows[0].cells,
                           ["规则", "复算公式", "结论", "置信度", "信源"]):
            cell.text = t
        for r_ in s["fact_check"]:
            cells = ft.add_row().cells
            for cell, v in zip(cells, [r_.get("issue_rule", ""), r_.get("formula", ""),
                                        r_.get("verdict", ""), r_.get("confidence", ""),
                                        r_.get("source", "")]):
                cell.text = str(v)

    # 五、批评修正记录
    if s["critic_notes"]:
        doc.add_heading("五、批评修正记录（误报压制）", level=1)
        for n in s["critic_notes"]:
            doc.add_paragraph(
                f"{n.get('issue_id')} → {n.get('action')}"
                f"（{n.get('by')}）：{n.get('reason')}", style="List Bullet")

    # 六、自动修订留痕
    rev = s["revision"]
    if rev.get("edits"):
        doc.add_heading("六、自动修订留痕", level=1)
        doc.add_paragraph(f"共自动应用 {rev.get('changed', 0)} 处修改。")
        rt = doc.add_table(rows=1, cols=4)
        rt.style = "Light Grid Accent 1"
        for cell, t in zip(rt.rows[0].cells, ["位置", "原文", "改为", "规则"]):
            cell.text = t
        for e in rev["edits"]:
            cells = rt.add_row().cells
            for cell, v in zip(cells, [e.get("span", ["", ""])[0], e.get("before", ""),
                                        e.get("after", ""), e.get("rule_id", "")]):
                cell.text = str(v)

    # 七、Agent 执行轨迹
    if s["trace"]:
        doc.add_heading("七、Agent 执行轨迹", level=1)
        tt = doc.add_table(rows=1, cols=5)
        tt.style = "Light Grid Accent 1"
        for cell, t in zip(tt.rows[0].cells,
                           ["时间", "节点", "智能体", "摘要", "耗时(ms)"]):
            cell.text = t
        for t in s["trace"]:
            cells = tt.add_row().cells
            for cell, v in zip(cells, [t.get("ts", ""), t.get("node", ""),
                                        t.get("agent", ""), t.get("summary", ""),
                                        t.get("elapsed_ms", "-")]):
                cell.text = str(v)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ======================================================================
# PDF (reportlab)
# ======================================================================
def export_pdf(report: AuditReport, text: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (ListFlowable, ListItem, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    FONT = "STSong-Light"

    ss = getSampleStyleSheet()
    base = ParagraphStyle("cn", parent=ss["Normal"], fontName=FONT,
                          fontSize=9.5, leading=14)
    h0 = ParagraphStyle("h0", parent=ss["Title"], fontName=FONT,
                        fontSize=17, leading=22, textColor=colors.HexColor("#0A1F44"))
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName=FONT,
                        fontSize=13, leading=18, textColor=colors.HexColor("#0A1F44"),
                        spaceBefore=10, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName=FONT,
                        fontSize=11, leading=15)
    para = ParagraphStyle("para", parent=base, spaceAfter=4)

    def P(text: Any, style=base) -> Paragraph:
        return Paragraph(str(text).replace("&", "&amp;").replace("<", "&lt;")
                         .replace(">", "&gt;"), style)

    s = _sections(report)
    gate_color, gate_label, gate_desc = s["gate"]

    elems: List[Any] = []
    elems.append(P(f"智能审校报告｜{report.doc_title}", h0))
    elems.append(P(f"签发门禁：{gate_label}（{report.release_gate}）　"
                   f"综合评分 {report.overall_score}/100　文体：{report.doc_type}"
                   f"　风格预设：{report.style_preset or '未指定'}"))
    elems.append(P(
        f"字数 {report.char_count}｜段落 {report.paragraph_count}｜"
        f"术语库版本 {report.terminology_version}｜推理后端 {report.llm_backend}"
        f"{'（已降级为规则引擎）' if report.degraded else ''}"))
    elems.append(Spacer(1, 6))

    # 一、维度评分
    elems.append(P("一、维度评分", h1))
    dim_data = [[P("维度"), P("得分"), P("致命"), P("严重"), P("一般"),
                 P("提示"), P("合计")]]
    for ds in s["dimensions"]:
        dim_data.append([P(ds.dimension.value), P(ds.score), P(ds.fatal),
                         P(ds.major), P(ds.minor), P(ds.info), P(ds.total)])
    dim_data.append([P("综合"), P(report.overall_score), P(s["counts"]["FATAL"]),
                     P(s["counts"]["MAJOR"]), P(s["counts"]["MINOR"]),
                     P(s["counts"]["INFO"]), P(len(report.issues))])
    t = Table(dim_data, colWidths=[92, 36, 32, 32, 32, 32, 36])
    t.setStyle(_table_style())
    elems.append(t)

    # 二、审校结论
    elems.append(P("二、审校结论", h1))
    elems.append(P(report.summary or "（无）", para))

    # 三、问题清单
    elems.append(P("三、问题清单", h1))
    for cat, items in s["grouped"].items():
        elems.append(P(f"{cat}（{len(items)} 条）", h2))
        for idx, issue in enumerate(items, 1):
            lines = _issue_lines(issue)
            sub = [P(f"{idx}. {lines[0]}",
                     ParagraphStyle("il", parent=base, leftIndent=8,
                                    spaceAfter=1))]
            for ln in lines[1:]:
                sub.append(P(ln, ParagraphStyle("is", parent=base,
                                               leftIndent=18, textColor=colors.HexColor("#555"))))
            elems.append(ListFlowable([ListItem(f, leftIndent=6, value="")
                                       for f in sub], bulletType="bullet",
                                      start="square"))
        elems.append(Spacer(1, 4))

    # 四、复算
    if s["fact_check"]:
        elems.append(P("四、数据与事实复算记录", h1))
        fc = [[P("规则"), P("复算公式"), P("结论"), P("置信度"), P("信源")]]
        for r_ in s["fact_check"]:
            fc.append([P(r_.get("issue_rule", "")), P(r_.get("formula", "")),
                       P(r_.get("verdict", "")), P(r_.get("confidence", "")),
                       P(r_.get("source", ""))])
        t = Table(fc, colWidths=[60, 130, 70, 50, 90])
        t.setStyle(_table_style())
        elems.append(t)

    # 五、批评修正
    if s["critic_notes"]:
        elems.append(P("五、批评修正记录（误报压制）", h1))
        items = [ListItem(P(f"{n.get('issue_id')} → {n.get('action')}"
                            f"（{n.get('by')}）：{n.get('reason')}"),
                          leftIndent=10) for n in s["critic_notes"]]
        elems.append(ListFlowable(items, bulletType="bullet"))

    # 六、修订留痕
    rev = s["revision"]
    if rev.get("edits"):
        elems.append(P("六、自动修订留痕", h1))
        elems.append(P(f"共自动应用 {rev.get('changed', 0)} 处修改。"))
        rd = [[P("位置"), P("原文"), P("改为"), P("规则")]]
        for e in rev["edits"]:
            span = e.get("span", ["", ""])
            rd.append([P(span[0]), P(e.get("before", "")), P(e.get("after", "")),
                       P(e.get("rule_id", ""))])
        t = Table(rd, colWidths=[40, 150, 150, 90])
        t.setStyle(_table_style())
        elems.append(t)

    # 七、轨迹
    if s["trace"]:
        elems.append(P("七、Agent 执行轨迹", h1))
        td = [[P("时间"), P("节点"), P("智能体"), P("摘要"), P("耗时")]]
        for tr in s["trace"]:
            td.append([P(tr.get("ts", "")), P(tr.get("node", "")),
                       P(tr.get("agent", "")), P(tr.get("summary", "")),
                       P(tr.get("elapsed_ms", "-"))])
        t = Table(td, colWidths=[80, 70, 70, 150, 50])
        t.setStyle(_table_style())
        elems.append(t)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=f"智能审校报告_{report.doc_title}",
                            leftMargin=18, rightMargin=18, topMargin=16, bottomMargin=16)
    doc.build(elems)
    return buf.getvalue()


def _table_style() -> "Any":
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle
    return TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A1F44")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C2D8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F2F5FB")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])
